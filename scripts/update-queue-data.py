"""Download and retain CAISO queue editions; never substitute zero for missing MW."""
import argparse
import datetime as dt
import hashlib
import io
import json
import math
from pathlib import Path
import re
import subprocess
import urllib.request

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
PAGE = "https://www.caiso.com/generation-transmission/generation/generator-interconnection"
URLS = {
  "legacy": "https://www.caiso.com/documents/publicqueuereport.xlsx",
  "cluster15": "https://www.caiso.com/documents/cluster-15-interconnection-requests.xlsx",
}


def text(value):
  return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def date(value):
  if value is None or text(value).upper() in ("", "N/A", "NA", "TBD"):
    return None
  if isinstance(value, (dt.datetime, dt.date)):
    return value.strftime("%Y-%m-%d")
  for fmt in ("%m/%d/%Y", "%Y-%m-%d"):
    try:
      return dt.datetime.strptime(text(value), fmt).strftime("%Y-%m-%d")
    except ValueError:
      pass
  raise ValueError(f"Invalid date: {value!r}")


def mw(value):
  if value is None or text(value).upper() in ("", "N/A", "NA"):
    return None
  result = float(value)
  if not math.isfinite(result) or result < 0:
    raise ValueError(f"Invalid MW: {value!r}")
  return result


def read_workbook(data, source_id, report_date=None):
  workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
  projects = []
  expected = (["Grid GenerationQueue", "Completed Generation Projects", "Withdrawn Generation Projects"]
              if source_id == "legacy" else ["Cluster 15 ", "Withdrawn"])
  if workbook.sheetnames != expected:
    raise ValueError(f"Unexpected worksheets: {workbook.sheetnames}")
  for sheet in workbook:
    rows = list(sheet.values)
    legacy = source_id == "legacy"
    if legacy:
      dates = [date(text(v).split(": ")[-1]) for v in rows[0] if text(v).startswith("Report Run Date:")]
      if len(dates) != 1 or (report_date and dates[0] != report_date):
        raise ValueError("Missing or inconsistent report run date")
      report_date = dates[0]
    header_index = 3 if legacy else 0
    headers = [text(v) for v in rows[header_index]]
    required = (["Queue Position", "Queue Date", "Application Status", "Net MWs to Grid", "Station or Transmission Line", "County", "State", "Utility", "Study Process", "Interconnection Request Receive Date", "Proposed On-line Date (as filed with IR)", "Interconnection Agreement Status", "Full Capacity, Partial or Energy Only (FC/P/EO)"] if legacy else
                ["Queue Number", "Queue Date", "Project Name", "NET MW POI", "POI", "PROJECT COUNTY", "Project State", "PTO", "Requested COD", "Application Date", "Study Area", "Service Type"])
    required += [f"Fuel-{i}" if legacy else f"Generation/Fuel {i}" for i in range(1, 4)]
    required += [f"MW-{i}" if legacy else f"NET MW {i}" for i in range(1, 4)]
    if legacy:
      required += ["Actual On-line Date" if "Completed" in sheet.title else "Current On-line Date"]
    if "Withdrawn" in sheet.title:
      required += ["Withdrawn Date" if legacy else "Withdrawal Date"]
    for name in required:
      if headers.count(name) != 1:
        raise ValueError(f"Missing or duplicate header: {name}")
    for row_number, values in enumerate(rows[header_index + 1:], header_index + 2):
      row = dict(zip(headers, values))
      queue_id = text(row.get("Queue Position" if legacy else "Queue Number"))
      if not queue_id:
        continue
      if not re.fullmatch(r"\d+[A-Za-z]*", queue_id):
        raise ValueError(f"Invalid queue identifier {queue_id}")
      status = text(row["Application Status"]) if legacy else ("WITHDRAWN" if sheet.title == "Withdrawn" else "ACTIVE")
      if status not in ("ACTIVE", "COMPLETED", "WITHDRAWN"):
        raise ValueError(f"Unexpected status {status}")
      components = []
      for i in range(1, 4):
        fuel = text(row.get(f"Fuel-{i}" if legacy else f"Generation/Fuel {i}"))
        capacity = mw(row.get(f"MW-{i}" if legacy else f"NET MW {i}"))
        if fuel and fuel.upper() not in ("N/A", "NA"):
          components.append({"fuel": fuel, "capacityMw": capacity})
        elif capacity is not None:
          components.append({"fuel": "Unspecified", "capacityMw": capacity})
      projects.append({
        "id": queue_id, "sourceId": source_id, "sourceSheet": sheet.title, "sourceRow": row_number,
        "name": text(row.get("Project Name", row.get("Project Name - Confidential"))) or f"Queue {queue_id}",
        "status": status, "queueDate": date(row["Queue Date"]),
        "applicationDate": date(row.get("Interconnection Request Receive Date" if legacy else "Application Date")),
        "withdrawnDate": date(row.get("Withdrawn Date" if legacy else "Withdrawal Date")),
        "completedDate": date(row.get("Actual On-line Date")),
        "originalOnlineDate": date(row.get("Proposed On-line Date (as filed with IR)" if legacy else "Requested COD")),
        "currentOnlineDate": date(row.get("Current On-line Date" if legacy else "Requested COD")),
        "netMw": mw(row["Net MWs to Grid" if legacy else "NET MW POI"]), "components": components,
        "county": text(row["County" if legacy else "PROJECT COUNTY"]),
        "state": text(row["State" if legacy else "Project State"]).upper(),
        "utility": text(row["Utility" if legacy else "PTO"]),
        "poi": text(row["Station or Transmission Line" if legacy else "POI"]),
        "studyProcess": text(row["Study Process"]) if legacy else "C15",
        "studyRegion": text(row.get("PTO Study Region" if legacy else "Study Area")),
        "agreementStatus": text(row.get("Interconnection Agreement Status")) or None,
        "deliverability": text(row.get("Full Capacity, Partial or Energy Only (FC/P/EO)" if legacy else "Service Type")) or None,
      })
  if not projects or not report_date:
    raise ValueError("Empty queue or missing report date")
  return projects, report_date


def fetch(url):
  with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "CaliforniaQueueExplorer/1.0"}), timeout=90) as response:
    return response.read()


def write_json(path, value):
  path.parent.mkdir(parents=True, exist_ok=True)
  temporary = path.with_suffix(".tmp")
  temporary.write_text(json.dumps(value, ensure_ascii=False, separators=(",", ":")) + "\n")
  temporary.replace(path)


def update(directory, inputs=None, force=False):
  now = dt.datetime.now(dt.timezone.utc)
  index_path = directory / "index.json"
  index = json.loads(index_path.read_text()) if index_path.exists() else {"schemaVersion": 1, "snapshots": []}
  if not force and index.get("checkedAt") and (now - dt.datetime.fromisoformat(index["checkedAt"])).total_seconds() < 7 * 86400:
    print("Queue sources checked within seven days; retaining current editions.")
    return
  html = (inputs / "caiso-interconnection.html").read_text() if inputs else fetch(PAGE).decode()
  match = re.search(r'cluster-15-interconnection-requests\.xlsx[\s\S]*?<span class="time">(\d{2}/\d{2}/\d{4})', html)
  if not match:
    raise ValueError("Cannot establish Cluster 15 publication date")
  projects, sources, raw_sources = [], [], []
  for source_id, url in URLS.items():
    filename = "caiso-queue.xlsx" if source_id == "legacy" else "caiso-cluster15.xlsx"
    data = (inputs / filename).read_bytes() if inputs else fetch(url)
    records, report_date = read_workbook(data, source_id, date(match[1]) if source_id == "cluster15" else None)
    digest = hashlib.sha256(data).hexdigest()
    source_path = f"sources/{source_id}-{digest[:16]}.xlsx"
    sources.append({"id": source_id, "url": url, "reportDate": report_date, "dateKind": "run date" if source_id == "legacy" else "publication date", "sha256": digest, "file": source_path})
    raw_sources.append((source_path, data))
    projects.extend(records)
  ids = [p["id"] for p in projects]
  if len(set(ids)) != len(ids):
    raise ValueError("Overlapping queue identifiers across reports; review source coverage before merging")
  projects.sort(key=lambda p: p["id"])
  # Changes in a workbook's run date or binary formatting alone do not create project-change events.
  fingerprint = hashlib.sha256(json.dumps(projects, sort_keys=True).encode()).hexdigest()[:16]
  observation_hash = hashlib.sha256((fingerprint + now.isoformat()).encode()).hexdigest()[:16]
  snapshot_id = now.strftime("%Y-%m-%d") + "-" + observation_hash
  previous = None
  if index["snapshots"]:
    previous = json.loads((directory / index["snapshots"][-1]["file"]).read_text())
    for source in sources:
      prior = next(s for s in index["latestSources"] if s["id"] == source["id"])
      if source["reportDate"] < prior["reportDate"]:
        raise ValueError("Source edition moved backward; retaining saved data")
  snapshot = {"schemaVersion": 1, "id": snapshot_id, "fingerprint": fingerprint, "collectedAt": now.isoformat(), "sources": sources, "projects": projects}
  validation_module = (ROOT / "assets/js/queue-data.js").as_uri()
  subprocess.run(["node", "--input-type=module", "-e", f'import fs from "node:fs"; import {{validateQueueSnapshot}} from "{validation_module}"; validateQueueSnapshot(JSON.parse(fs.readFileSync(0,"utf8")));'], input=json.dumps(snapshot), text=True, check=True)
  for source_path, data in raw_sources:
    target = directory / source_path
    target.parent.mkdir(parents=True, exist_ok=True)
    if not target.exists():
      target.write_bytes(data)
  if not previous or previous["fingerprint"] != fingerprint:
    relative = f"snapshots/{snapshot_id}.json"
    write_json(directory / relative, snapshot)
    index["snapshots"].append({"id": snapshot_id, "file": relative, "collectedAt": now.isoformat()})
  index["checkedAt"] = now.isoformat()
  index["latestSources"] = sources
  write_json(index_path, index)
  print(f"Verified {len(projects)} projects; {sum(p['status'] == 'ACTIVE' for p in projects)} active; {len(index['snapshots'])} snapshot(s).")


if __name__ == "__main__":
  parser = argparse.ArgumentParser()
  parser.add_argument("--directory", type=Path, default=ROOT / "assets/data/queue")
  parser.add_argument("--inputs", type=Path, help="Directory containing the two source workbooks and source listing HTML")
  parser.add_argument("--force", action="store_true")
  args = parser.parse_args()
  update(args.directory, args.inputs, args.force)
